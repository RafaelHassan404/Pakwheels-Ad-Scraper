import sqlite3
import json

DB_FILE = "pakwheels.db"

def init_db():
    conn = sqlite3.connect(DB_FILE)
    cur = conn.cursor()

    cur.execute("""
        CREATE TABLE IF NOT EXISTS listings (
            ad_id INTEGER PRIMARY KEY,
            brand TEXT,
            model TEXT,
            year INTEGER,
            mileage TEXT,
            engine_type TEXT,
            location TEXT,
            link TEXT,
            image_url TEXT,
            description TEXT,
            pros TEXT,
            cons TEXT,
            deal_score REAL,
            current_price INTEGER,
            status TEXT DEFAULT 'active',
            first_seen TEXT,
            last_seen TEXT,
            sold_detected_date TEXT
        )
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS price_history (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            ad_id INTEGER,
            price INTEGER,
            recorded_date TEXT,
            FOREIGN KEY (ad_id) REFERENCES listings(ad_id)
        )
    """)

    conn.commit()
    conn.close()
    print("Database initialized.")

if __name__ == "__main__":
    init_db()

from datetime import date

def save_listings(listings):
    if not listings:
        print("No listings passed in — skipping save.")
        return []

    conn = sqlite3.connect(DB_FILE)
    cur = conn.cursor()
    today = date.today().isoformat()

    seen_ad_ids = []

    for item in listings:
        ad_id = int(item["Link"].rstrip("/").split("-")[-1])
        seen_ad_ids.append(ad_id)

        cur.execute("SELECT current_price FROM listings WHERE ad_id = ?", (ad_id,))
        existing = cur.fetchone()

        if existing is None:
            cur.execute("""
                INSERT INTO listings (ad_id, brand, model, year, mileage, engine_type,
                    location, link, image_url, description, current_price, status, first_seen, last_seen)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'active', ?, ?)
            """, (ad_id, item["Brand"], item["Model"], item["Year"], item["Mileage"],
                  item["Engine_Type"], item["Location"], item["Link"], item["Image_URL"],
                  item.get("Description", ""), item["Price_PKR"], today, today))
            cur.execute("INSERT INTO price_history (ad_id, price, recorded_date) VALUES (?, ?, ?)",
                        (ad_id, item["Price_PKR"], today))
        else:
            old_price = existing[0]
            if item.get("Description") is not None:
                cur.execute("""
                    UPDATE listings SET current_price = ?, status = 'active', last_seen = ?,
                        mileage = ?, location = ?, description = ?
                    WHERE ad_id = ?
                """, (item["Price_PKR"], today, item["Mileage"], item["Location"],
                    item["Description"], ad_id))
            else:
                cur.execute("""
                    UPDATE listings SET current_price = ?, status = 'active', last_seen = ?,
                        mileage = ?, location = ?
                    WHERE ad_id = ?
                """, (item["Price_PKR"], today, item["Mileage"], item["Location"], ad_id))

            if old_price != item["Price_PKR"]:
                cur.execute("INSERT INTO price_history (ad_id, price, recorded_date) VALUES (?, ?, ?)",
                            (ad_id, item["Price_PKR"], today))

    conn.commit()
    conn.close()
    print(f"Saved {len(listings)} listings.")
    return seen_ad_ids


def mark_sold_listings(seen_ad_ids):
    if not seen_ad_ids:
        print("No ad_ids passed in — skipping sold detection to avoid marking everything as sold.")
        return

    conn = sqlite3.connect(DB_FILE)
    cur = conn.cursor()
    today = date.today().isoformat()

    placeholders = ",".join("?" * len(seen_ad_ids))
    cur.execute(f"""
        UPDATE listings SET status = 'sold', sold_detected_date = ?
        WHERE status = 'active' AND ad_id NOT IN ({placeholders})
    """, (today, *seen_ad_ids))

    conn.commit()
    updated = cur.rowcount
    conn.close()
    print(f"{updated} listings marked as sold.")


def get_sold_listings():
    conn = sqlite3.connect(DB_FILE)
    cur = conn.cursor()
    cur.execute("""
        SELECT ad_id, brand, model, year, current_price, sold_detected_date
        FROM listings
        WHERE status = 'sold'
        ORDER BY sold_detected_date DESC
    """)
    rows = cur.fetchall()
    conn.close()
    return rows

def get_listings_needing_analysis():
    conn = sqlite3.connect(DB_FILE)
    cur = conn.cursor()
    cur.execute("""
        SELECT ad_id, description FROM listings
        WHERE description IS NOT NULL AND description != ''
        AND (pros IS NULL OR cons IS NULL)
    """)
    rows = cur.fetchall()
    conn.close()
    return rows

def update_pros_cons(ad_id, pros, cons):
    conn = sqlite3.connect(DB_FILE)
    cur = conn.cursor()
    cur.execute("UPDATE listings SET pros = ?, cons = ? WHERE ad_id = ?", (pros, cons, ad_id))
    conn.commit()
    conn.close()

def init_market_values():
    conn = sqlite3.connect(DB_FILE)
    cur = conn.cursor()
    cur.execute("""
        CREATE TABLE IF NOT EXISTS market_values (
            brand TEXT,
            model TEXT,
            year TEXT,
            avg_price REAL,
            median_price REAL,
            sample_size INTEGER,
            last_updated TEXT,
            PRIMARY KEY (brand, model, year)
        )
    """)
    conn.commit()
    conn.close()

def refresh_market_values():
    conn = sqlite3.connect(DB_FILE)
    cur = conn.cursor()

    cur.execute("""
        SELECT brand, model, year, current_price
        FROM listings
        WHERE status = 'active' AND current_price IS NOT NULL
    """)
    rows = cur.fetchall()

    groups = {}
    for brand, model, year, price in rows:
        key = (brand, model, year)
        groups.setdefault(key, []).append(price)

    today = date.today().isoformat()
    for (brand, model, year), prices in groups.items():
        prices.sort()
        avg_price = sum(prices) / len(prices)
        n = len(prices)
        median_price = prices[n // 2] if n % 2 else (prices[n // 2 - 1] + prices[n // 2]) / 2

        cur.execute("""
            INSERT INTO market_values (brand, model, year, avg_price, median_price, sample_size, last_updated)
            VALUES (?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(brand, model, year) DO UPDATE SET
                avg_price = excluded.avg_price,
                median_price = excluded.median_price,
                sample_size = excluded.sample_size,
                last_updated = excluded.last_updated
        """, (brand, model, year, avg_price, median_price, n, today))

    conn.commit()
    conn.close()
    print(f"Market values updated for {len(groups)} brand/model/year combinations.")

def refresh_deal_scores():
    conn = sqlite3.connect(DB_FILE)
    cur = conn.cursor()
    cur.execute("""
        SELECT l.ad_id, l.current_price, m.median_price
        FROM listings l
        JOIN market_values m ON l.brand = m.brand AND l.model = m.model AND l.year = m.year
        WHERE l.status = 'active' AND m.sample_size >= 3
    """)
    rows = cur.fetchall()

    for ad_id, price, median in rows:
        score = round((median - price) / median * 100, 1)  # % below market, positive = good deal
        cur.execute("UPDATE listings SET deal_score = ? WHERE ad_id = ?", (score, ad_id))

    conn.commit()
    conn.close()
    print(f"Deal scores updated for {len(rows)} listings.")

def get_existing_ad_ids():
    conn = sqlite3.connect(DB_FILE)
    cur = conn.cursor()
    cur.execute("SELECT ad_id FROM listings WHERE description IS NOT NULL AND description != ''")
    ids = {row[0] for row in cur.fetchall()}
    conn.close()
    return ids

def export_to_excel(filepath="pakwheels_report.xlsx"):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter

    conn = sqlite3.connect(DB_FILE)
    cur = conn.cursor()
    cur.execute("""
        SELECT ad_id, brand, model, year, mileage, engine_type, location,
               current_price, deal_score, pros, cons, status, link, first_seen, last_seen
        FROM listings
        ORDER BY deal_score DESC
    """)
    rows = cur.fetchall()
    conn.close()

    if not rows:
        print("No listings to export.")
        return

    headers = ["Ad ID", "Brand", "Model", "Year", "Mileage", "Engine", "Location",
               "Price", "Deal Score (%)", "Pros", "Cons", "Status", "Link",
               "First Seen", "Last Seen"]

    wb = Workbook()
    sheet = wb.active
    sheet.title = "Bike Listings"

    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    good_deal_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    sold_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

    for row in rows:
        ad_id, brand, model, year, mileage, engine, location, price, deal_score, pros, cons, status, link, first_seen, last_seen = row

        pros_str = ", ".join(json.loads(pros)) if pros else ""
        cons_str = ", ".join(json.loads(cons)) if cons else ""

        sheet.append([ad_id, brand, model, year, mileage, engine, location,
                       price, deal_score, pros_str, cons_str, status, link, first_seen, last_seen])

        current_row = sheet.max_row
        if status == "sold":
            for cell in sheet[current_row]:
                cell.fill = sold_fill
        elif deal_score is not None and deal_score >= 10:
            for cell in sheet[current_row]:
                cell.fill = good_deal_fill

    for i, header in enumerate(headers, 1):
        max_len = max(
            [len(str(header))] + [len(str(sheet.cell(row=r, column=i).value or "")) for r in range(2, sheet.max_row + 1)]
        )
        sheet.column_dimensions[get_column_letter(i)].width = min(max_len + 3, 50)

    sheet.freeze_panes = "A2"

    price_col = headers.index("Price") + 1
    for r in range(2, sheet.max_row + 1):
        sheet.cell(row=r, column=price_col).number_format = '#,##0'

    wb.save(filepath)
    print(f"Exported {len(rows)} listings to {filepath}")